import openpyxl
import re

# Load Excel
excel_path = r"src/assets/FCI Stock28092026.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)
sheet = wb['FCI STOCK 28-09-2026']
excel_products = []
for r in range(2, sheet.max_row + 1):
    name = sheet.cell(row=r, column=1).value
    qty = sheet.cell(row=r, column=2).value
    rate = sheet.cell(row=r, column=3).value
    if name is None and qty is None and rate is None:
        continue
    excel_products.append(str(name).strip() if name else '')
print(f"Excel total products: {len(excel_products)}")

# Load products.ts
content = open("src/data/products.ts", encoding="utf-8").read()

# Extract sku values
skus = re.findall(r'sku:\s*["\']([^"\']+)["\']', content)
print(f"Products.ts total sku entries: {len(skus)}")

# Check how many excel product names appear anywhere in the TS file
found_in_ts = [ep for ep in excel_products if ep in content]
not_in_ts = [ep for ep in excel_products if ep not in content]

print(f"\nExcel product names found in products.ts: {len(found_in_ts)} / {len(excel_products)}")
print(f"Excel product names NOT in products.ts: {len(not_in_ts)}")
print()
print("First 20 NOT in products.ts:")
for ep in not_in_ts[:20]:
    print(f"  - {ep}")
