import openpyxl
import json

excel_path = r"c:\My Drive\Freelance\Projects\qualitech-landing-page\src\assets\FCI Stock28092026.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)

print("Sheets in workbook:", wb.sheetnames)

for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    print(f"\n--- Sheet: {sheet_name} (max_row={sheet.max_row}, max_column={sheet.max_column}) ---")
    
    # Read first 15 rows to understand the structure / header row
    for r in range(1, min(16, sheet.max_row + 1)):
        row_vals = [sheet.cell(row=r, column=c).value for c in range(1, min(25, sheet.max_column + 1))]
        if any(v is not None for v in row_vals):
            print(f"Row {r}: {row_vals}")

