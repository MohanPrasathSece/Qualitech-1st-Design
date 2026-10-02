import openpyxl
import re
import json

excel_path = r"c:\My Drive\Freelance\Projects\qualitech-landing-page\src\assets\FCI Stock28092026.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)
sheet = wb['FCI STOCK 28-09-2026']

products = []

for r in range(2, sheet.max_row + 1):
    raw_name = sheet.cell(row=r, column=1).value
    qty = sheet.cell(row=r, column=2).value
    rate = sheet.cell(row=r, column=3).value
    note = sheet.cell(row=r, column=4).value
    
    if raw_name is None and qty is None and rate is None:
        continue
        
    raw_name_str = str(raw_name).strip() if raw_name is not None else ""
    qty_num = int(qty) if isinstance(qty, (int, float)) else 0
    rate_num = float(rate) if isinstance(rate, (int, float)) else 0.0
    
    products.append({
        'row': r,
        'raw_name': raw_name_str,
        'qty': qty_num,
        'rate': rate_num,
        'note': str(note).strip() if note is not None else ""
    })

print(f"Total parsed rows: {len(products)}")

# Let's inspect naming patterns and categories
# FCI / Amphenol product series patterns:
# e.g., 10067847-..., 10127815-..., 8604..., 8613..., 8655..., BergStik, Quickie, Minitek, Millipacs, D-Sub, RJ45, D-Sub connectors (SEMAT, SDMAT, etc.), OFC Patch Cables (SC-SC...), etc.

categories = {}
series_patterns = []

for p in products:
    name = p['raw_name']
    
    # Check for fiber patch cables
    if 'OPTICAL FIBER' in name.upper() or 'PATCH CABLE' in name.upper() or name.startswith('SC-') or name.startswith('LC-'):
        cat = "Fiber Optic Cables"
        sub = "OFC Patch Cables"
    elif name.startswith('8655') or 'D-SUB' in name.upper() or name.startswith('SEMA') or name.startswith('SDMA') or name.startswith('SF') or name.startswith('SS'):
        cat = "Connectors"
        sub = "D-Subminiature & Backshells"
    elif name.startswith('10127815') or name.startswith('10127816') or name.startswith('10127817') or name.startswith('10127716') or name.startswith('10127718'):
        cat = "Connectors"
        sub = "Wire / Cable / Flex to Board"
    elif name.startswith('10131936') or name.startswith('10132450') or name.startswith('10135551') or name.startswith('10135554') or name.startswith('10138980'):
        cat = "Connectors"
        sub = "Board to Board"
    elif name.startswith('8613') or name.startswith('8614') or name.startswith('8604') or name.startswith('719') or name.startswith('773') or name.startswith('680') or name.startswith('679') or name.startswith('686'):
        cat = "Connectors"
        sub = "Headers & Receptacles"
    elif name.startswith('100') or name.startswith('101'):
        cat = "Connectors"
        sub = "Board to Board / High Density"
    else:
        cat = "Connectors"
        sub = "Interconnect Systems"
        
    categories[f"{cat} > {sub}"] = categories.get(f"{cat} > {sub}", 0) + 1

print("\nCategory Distribution:")
for k, v in sorted(categories.items()):
    print(f"  {k}: {v}")

# Price summary
prices = [p['rate'] for p in products if p['rate'] > 0]
print(f"\nPrice stats (non-zero prices: {len(prices)} / {len(products)}):")
print(f"  Min: ₹{min(prices):.2f}")
print(f"  Max: ₹{max(prices):.2f}")
print(f"  Avg: ₹{sum(prices)/len(prices):.2f}")

# Stock summary
stocks = [p['qty'] for p in products]
print(f"\nStock stats:")
print(f"  In-stock (>0): {sum(1 for s in stocks if s > 0)}")
print(f"  Out-of-stock (0): {sum(1 for s in stocks if s <= 0)}")
print(f"  Total units in stock: {sum(stocks)}")

