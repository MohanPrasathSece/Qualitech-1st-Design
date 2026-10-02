import openpyxl
import requests
import json
import re
import math

SUPABASE_URL = "https://xgttkvykqtoppbsfukzr.supabase.co"
SERVICE_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InhndHRrdnlrcXRvcHBic2Z1a3pyIiwicm9sZSI6InNlcnZpY2Vfcm9sZSIsImlhdCI6MTc4OTI3MjcxNywiZXhwIjoyMTA0ODQ4NzE3fQ.iWnF2ZuGuO439mBXfWtLESKw7sxzTFDQGNkyhKZNFUs"

HEADERS = {
    "apikey": SERVICE_KEY,
    "Authorization": f"Bearer {SERVICE_KEY}",
    "Content-Type": "application/json",
}

def categorize(name):
    n = name.upper()
    if "OPTICAL FIBER" in n or "PATCH CABLE" in n or name.startswith("SC-") or name.startswith("LC-"):
        return "Fiber Optic Cables", "OFC Patch Cables"
    if name.startswith("8655") or "D-SUB" in n or name.startswith("SEMA") or name.startswith("SDMA") or name.startswith("SEMAT") or name.startswith("SDMAT") or name.startswith("SF") or name.startswith("SS"):
        return "Connectors", "D-Subminiature & Backshells"
    if name.startswith("10127715") or name.startswith("10127716") or name.startswith("10127717") or name.startswith("10127718") or name.startswith("10127815") or name.startswith("10127816") or name.startswith("10127817"):
        return "Connectors", "Wire to Board"
    if name.startswith("10131936") or name.startswith("10132450") or name.startswith("10135551") or name.startswith("10135554") or name.startswith("10138980"):
        return "Connectors", "Board to Board"
    if name.startswith("8613") or name.startswith("8614") or name.startswith("8604") or name.startswith("8606") or name.startswith("719") or name.startswith("773") or name.startswith("680") or name.startswith("679") or name.startswith("686"):
        return "Connectors", "Headers & Receptacles"
    if name.startswith("100") or name.startswith("101"):
        return "Connectors", "Board to Board"
    return "Connectors", "Interconnect Systems"

def slugify(name):
    slug = re.sub(r"[^a-zA-Z0-9\-]", "-", name)
    slug = re.sub(r"-+", "-", slug).strip("-").lower()
    return f"fci-{slug[:60]}"

def safe_float(val):
    try:
        f = float(val)
        return 0.0 if (math.isnan(f) or math.isinf(f)) else round(f, 4)
    except:
        return 0.0

def safe_int(val):
    try:
        f = float(val)
        return 0 if (math.isnan(f) or math.isinf(f)) else int(f)
    except:
        return 0

print("Step 1: Deleting all existing products...")
r = requests.delete(f"{SUPABASE_URL}/rest/v1/products?id=neq.NOOP", headers={**HEADERS, "Prefer": "return=minimal"})
if r.status_code in (200, 204):
    print(f"  OK deleted (status {r.status_code})")
else:
    print(f"  FAIL: {r.status_code} {r.text}")
    exit(1)

print("\nStep 2: Parsing Excel...")
wb = openpyxl.load_workbook(r"src/assets/FCI Stock28092026.xlsx", data_only=True)
sheet = wb["FCI STOCK 28-09-2026"]
products = []
seen_ids = set()
for r_idx in range(2, sheet.max_row + 1):
    raw_name = sheet.cell(row=r_idx, column=1).value
    qty = sheet.cell(row=r_idx, column=2).value
    rate = sheet.cell(row=r_idx, column=3).value
    note = sheet.cell(row=r_idx, column=4).value
    if raw_name is None and qty is None and rate is None:
        continue
    name_str = str(raw_name).strip() if raw_name is not None else ""
    if not name_str:
        continue
    qty_num = safe_int(qty)
    rate_num = safe_float(rate)
    note_str = str(note).strip() if note else ""
    category, sub_category = categorize(name_str)
    prod_id = slugify(name_str)
    if prod_id in seen_ids:
        prod_id = f"{prod_id}-{r_idx}"
    seen_ids.add(prod_id)
    products.append({
        "id": prod_id, "sku": name_str, "name": name_str,
        "brand": "FCI / Amphenol", "category": category, "sub_category": sub_category,
        "description": note_str if note_str else f"{name_str} - FCI Amphenol connector component.",
        "features": [], "specs": {}, "industries": ["Electronics", "Industrial", "Telecom"],
        "featured": False, "image": "/logo.png", "in_stock": qty_num > 0,
        "external_url": None, "is_enquiry": False,
        "price": rate_num if rate_num > 0 else 0.0, "sale_price": None,
        "stock_count": qty_num, "low_stock_threshold": 10, "min_order_qty": 1,
        "unit": "pcs", "rating": 0.0, "review_count": 0,
        "lead_time": "2-5 Business Days", "tiered_pricing": [], "datasheet_url": None,
    })
print(f"  Parsed {len(products)} products")

print("\nStep 3: Upserting into Supabase...")
BATCH = 50
success = 0
for i in range(0, len(products), BATCH):
    batch = products[i:i+BATCH]
    bn = i // BATCH + 1
    resp = requests.post(
        f"{SUPABASE_URL}/rest/v1/products",
        headers={**HEADERS, "Prefer": "resolution=merge-duplicates,return=minimal"},
        data=json.dumps(batch)
    )
    if resp.status_code in (200, 201):
        print(f"  Batch {bn}: OK ({len(batch)} products)")
        success += len(batch)
    else:
        print(f"  Batch {bn}: FAIL {resp.status_code} {resp.text[:200]}")

print("\nStep 4: Verifying...")
vr = requests.get(f"{SUPABASE_URL}/rest/v1/products?select=id", headers={**HEADERS, "Prefer": "count=exact"})
print(f"  Final count in DB: {vr.headers.get('Content-Range')}")
print(f"\nDONE! {success}/{len(products)} products seeded.")
