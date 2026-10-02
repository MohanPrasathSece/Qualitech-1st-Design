import openpyxl
import re
import os
import json

excel_path = r"c:\My Drive\Freelance\Projects\qualitech-landing-page\src\assets\FCI Stock28092026.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)
sheet = wb['FCI STOCK 28-09-2026']

headers = [sheet.cell(row=1, column=c).value for c in range(1, sheet.max_column + 1)]
print("Headers:", headers)

rows = []
for r in range(2, sheet.max_row + 1):
    name_val = sheet.cell(row=r, column=1).value
    qty_val = sheet.cell(row=r, column=2).value
    rate_val = sheet.cell(row=r, column=3).value
    extra_val = sheet.cell(row=r, column=4).value
    
    if name_val is None and qty_val is None and rate_val is None:
        continue
    
    rows.append({
        'row_idx': r,
        'raw_name': name_val,
        'qty': qty_val,
        'rate': rate_val,
        'extra': extra_val
    })

print(f"Total non-empty data rows: {len(rows)}")

# Check for duplicates or formatting
names = [r['raw_name'] for r in rows]
print(f"Unique names: {len(set(names))}")

# Check any with extra column
with_extra = [r for r in rows if r['extra'] is not None]
print(f"Rows with extra col 4: {len(with_extra)}")
if with_extra:
    print("Sample with extra:", with_extra[:5])

# Check data types and ranges
null_names = [r for r in rows if not r['raw_name']]
invalid_qty = [r for r in rows if r['qty'] is None or not isinstance(r['qty'], (int, float))]
invalid_rate = [r for r in rows if r['rate'] is None or not isinstance(r['rate'], (int, float))]

print(f"Null names: {len(null_names)}")
print(f"Invalid qty: {len(invalid_qty)}")
print(f"Invalid rate: {len(invalid_rate)}")

qty_zero_count = sum(1 for r in rows if isinstance(r['qty'], (int, float)) and r['qty'] <= 0)
rate_zero_count = sum(1 for r in rows if isinstance(r['rate'], (int, float)) and r['rate'] <= 0)
print(f"Qty <= 0: {qty_zero_count}")
print(f"Rate <= 0: {rate_zero_count}")

# Look at first 20 rows
print("\nFirst 20 products:")
for r in rows[:20]:
    print(f"Row {r['row_idx']}: Name='{r['raw_name']}', Qty={r['qty']}, Rate={r['rate']}")

# Look at last 10 rows
print("\nLast 10 products:")
for r in rows[-10:]:
    print(f"Row {r['row_idx']}: Name='{r['raw_name']}', Qty={r['qty']}, Rate={r['rate']}")

